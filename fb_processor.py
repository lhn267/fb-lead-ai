import os
import json
import re
import time
import requests
import pandas as pd
from typing import List, Dict, Any, Tuple, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Common Vietnamese joke / fake workplace keywords
JOKE_KEYWORDS = [
    "chém gió", "tại nhà", "ở nhà", "ăn bám", "tâm thần", "thất nghiệp",
    "độc thân", "làm gì có việc", "bán muối", "đại học bôn ba", "học viện lừa đảo",
    "chơi game", "nằm ngủ", "nuôi lợn", "la cà", "trộm chó", "thất tình",
    "học sinh cá biệt", "bụi đời", "đang thất nghiệp"
]

def detect_columns(df: pd.DataFrame) -> Dict[str, Any]:
    """
    Intelligently inspects actual cell contents (just like ChatGPT) to classify columns into:
    - link_col: Facebook profile link
    - name_col: Person's name (Vietnamese / Latin names, excluding image URLs and descriptions)
    - info_cols: All description/job/education/mutual friend columns
    - img_cols: Avatar image URLs (scontent, fbcdn, etc.)
    """
    cols = df.columns.tolist()
    link_col = None
    name_col = None
    info_cols = []
    img_cols = []

    # Step 1: Detect Link and Image columns by examining content
    for col in cols:
        sample_vals = [str(x).strip() for x in df[col].dropna().head(15)]
        if not sample_vals:
            continue
        
        # Check if column contains image URLs
        if any(("scontent" in v or "fbcdn.net" in v or ".jpg" in v or ".png" in v or "data:image" in v) for v in sample_vals) or "src" in str(col).lower():
            img_cols.append(col)
            continue
            
        # Check if column is a Facebook profile link
        if any("facebook.com" in v for v in sample_vals) and not any("friends_mutual" in v for v in sample_vals):
            if not link_col:
                link_col = col
                continue

    # Step 2: From remaining columns, identify Name vs Info
    non_link_img = [c for c in cols if c != link_col and c not in img_cols and "href" not in str(c).lower()]
    
    # Priority for name: Check header names first if present
    for c in non_link_img:
        c_low = str(c).lower().strip()
        if any(k == c_low or k in c_low for k in ["name", "tên", "họ tên", "fullname", "full_name"]):
            name_col = c
            break

    # If no header match for name, inspect contents
    if not name_col:
        for c in non_link_img:
            sample_vals = [str(x).strip() for x in df[c].dropna().head(15)]
            if not sample_vals:
                continue
            # A person's name typically has:
            # - Short length (1 to 5 words)
            # - Does NOT contain keywords like "tại", "làm việc", "bạn chung", "sống tại", "trường", "đại học", "theo dõi", "http"
            is_candidate_name = all(
                len(v.split()) <= 5 and not any(kw in v.lower() for kw in ["tại", "làm việc", "bạn chung", "sống tại", "trường", "đại học", "theo dõi", "http", "giám đốc", "quản lý", "kinh doanh", "founder", "ceo", "chủ tịch", "nhân viên"])
                for v in sample_vals
            )
            if is_candidate_name:
                name_col = c
                break

    # Fallback for name if still not found
    if not name_col and non_link_img:
        name_col = non_link_img[0]

    # Step 3: All other non-image, non-link columns are info columns!
    for c in non_link_img:
        if c != name_col:
            info_cols.append(c)

    return {
        "name_col": name_col,
        "link_col": link_col,
        "info_cols": info_cols,
        "info_col": info_cols[0] if info_cols else (name_col or "")
    }

def normalize_fb_link(url: str) -> str:
    """Normalizes Facebook profile URL by removing tracking params for accurate deduplication."""
    if not url or not isinstance(url, str):
        return ""
    url = url.strip()
    if "profile.php" in url and "id=" in url:
        match = re.search(r"profile\.php\?id=(\d+)", url)
        if match:
            return f"https://www.facebook.com/profile.php?id={match.group(1)}"
    clean_url = url.split("?")[0].rstrip("/")
    return clean_url

def clean_file_data(uploaded_file) -> pd.DataFrame:
    """Reads uploaded CSV or Excel file safely with multiple encoding fallbacks."""
    file_name = getattr(uploaded_file, "name", str(uploaded_file))
    if file_name.endswith((".xlsx", ".xls")):
        return pd.read_excel(uploaded_file)
    else:
        # Try multiple encodings for CSV
        for enc in ["utf-8-sig", "utf-8", "cp1258", "latin1", "utf-16"]:
            try:
                if hasattr(uploaded_file, "seek"):
                    uploaded_file.seek(0)
                return pd.read_csv(uploaded_file, encoding=enc)
            except Exception:
                continue
        # Fallback
        if hasattr(uploaded_file, "seek"):
            uploaded_file.seek(0)
        return pd.read_csv(uploaded_file, encoding="utf-8", errors="replace")

def merge_and_deduplicate_dfs(dfs: List[pd.DataFrame]) -> Tuple[pd.DataFrame, int, int]:
    """
    Merges multiple DataFrames from different crawl sessions/batches and
    deduplicates based on Facebook link (normalized) or name.
    Returns: (deduplicated_df, total_rows_before, duplicates_removed_count)
    """
    if not dfs:
        return pd.DataFrame(), 0, 0
    valid_dfs = [d for d in dfs if d is not None and not d.empty]
    if not valid_dfs:
        return pd.DataFrame(), 0, 0

    combined = pd.concat(valid_dfs, ignore_index=True)
    total_before = len(combined)

    cols = detect_columns(combined)
    link_col = cols.get("link_col")
    name_col = cols.get("name_col")

    if link_col and link_col in combined.columns:
        temp_link_clean = combined[link_col].apply(normalize_fb_link)
        is_valid_link = (temp_link_clean != "") & (~combined[link_col].isna())

        valid_part = combined[is_valid_link].copy()
        valid_part["_norm_link"] = temp_link_clean[is_valid_link]
        valid_part = valid_part.drop_duplicates(subset=["_norm_link"], keep="first").drop(columns=["_norm_link"])

        invalid_part = combined[~is_valid_link]
        if name_col and name_col in invalid_part.columns:
            invalid_part = invalid_part.drop_duplicates(subset=[name_col], keep="first")

        deduped = pd.concat([valid_part, invalid_part], ignore_index=True)
    elif name_col and name_col in combined.columns:
        deduped = combined.drop_duplicates(subset=[name_col], keep="first")
    else:
        deduped = combined.drop_duplicates(keep="first")

    dupes_removed = total_before - len(deduped)
    return deduped, total_before, dupes_removed

def rule_based_pre_classify(text: str) -> Optional[Dict[str, Any]]:
    """
    Fast rule-based pre-screening to detect obvious jokes, empty descriptions, or mutual-friends only.
    """
    if not text or not isinstance(text, str) or not text.strip():
        return {
            "job_title": "Chưa cập nhật",
            "company": "Chưa cập nhật",
            "company_normalized": "Chưa cập nhật",
            "school": "Chưa cập nhật",
            "level": "Không xác định",
            "industry": "Không rõ",
            "status": "Không có thông tin việc làm",
            "lead_tier": "Bỏ qua",
            "reason": "Không có mô tả hoặc chỉ có tên"
        }

    text_lower = text.lower().strip()

    # Intercept profiles with no work/school info or Facebook's standard empty placeholders
    no_info_patterns = [
        "không có thông tin việc làm công khai",
        "no workplaces to show",
        "no schools/universities to show",
        "không có thông tin",
        "chưa cập nhật thông tin việc làm"
    ]
    if any(p in text_lower for p in no_info_patterns):
        # Verify if there is any other actual info in the text (like a bio or job title)
        stripped = text_lower
        for p in no_info_patterns:
            stripped = stripped.replace(p, "")
        stripped = re.sub(r"(tiểu sử/mô tả:|công việc:|hạng mục/lĩnh vực:|học vấn:|university|high school|\d+\s+bạn chung|\d+\s+mutual friends|[·|\n\r\t,])", " ", stripped).strip()
        if not stripped or len(stripped) < 3:
            return {
                "job_title": "Chưa cập nhật",
                "company": "Chưa cập nhật",
                "company_normalized": "Chưa cập nhật",
                "school": "Chưa cập nhật",
                "level": "Không xác định",
                "industry": "Không rõ",
                "status": "Không có thông tin việc làm",
                "lead_tier": "Bỏ qua",
                "reason": "Chưa cài đặt thông tin việc làm công khai trên Facebook"
            }

    # Check for joke keywords
    for kw in JOKE_KEYWORDS:
        if kw in text_lower:
            return {
                "job_title": "Ảo / Đùa cợt",
                "company": "Ảo",
                "company_normalized": "Ảo",
                "school": "Chưa cập nhật",
                "level": "Rác / Ảo",
                "industry": "Khác",
                "status": "Rác / Đùa cợt",
                "lead_tier": "Bỏ qua",
                "reason": f"Phát hiện từ khóa đùa cợt: '{kw}'"
            }

    # If only contains mutual friends e.g. "15 bạn chung" or "mutual friends"
    if re.search(r"^\s*\d+\s+bạn chung\s*$", text_lower) or re.search(r"^\s*\d+\s+mutual friends\s*$", text_lower):
        return {
            "job_title": "Chưa cập nhật",
            "company": "Chưa cập nhật",
            "company_normalized": "Chưa cập nhật",
            "school": "Chưa cập nhật",
            "level": "Không xác định",
            "industry": "Không rõ",
            "status": "Không có thông tin việc làm",
            "lead_tier": "Bỏ qua",
            "reason": "Chỉ hiển thị số bạn chung, không ghi công việc"
        }

    return None

def build_gemini_prompt(batch_items: List[Dict[str, Any]], target_criteria: str = "") -> str:
    """
    Builds structured prompt for AI batch extraction with separate School and Normalized Company fields.
    """
    prompt = f"""Bạn là một chuyên gia phân tích dữ liệu B2B và CRM, chuyên trích xuất thông tin khách hàng tiềm năng từ dữ liệu bạn bè Facebook.
Nhiệm vụ: Trích xuất chức vụ, tên công ty, chuẩn hóa tên công ty đầy đủ, trường học, cấp bậc, ngành nghề và đánh giá độ tiềm năng của danh sách bạn bè dưới đây.

{f"MỤC TIÊU KHÁCH HÀNG ƯU TIÊN (ICP): {target_criteria}" if target_criteria else "MỤC TIÊU: Tìm các chủ doanh nghiệp, C-Level, Quản lý, Chuyên viên có công việc rõ ràng, loại bỏ nick rác/đùa cợt."}

QUY TẮC BÓC TÁCH BẮT BUỘC:
1. `Tên công ty / Đơn vị (company)`:
   - Tên nơi làm việc, công ty, doanh nghiệp, cơ quan (VD: "GCW JSC", "iQIYI Vietnam", "VIONA Travel", "Đất Xanh Miền Bắc", "TechViet Solutions", "Vingroup", "FPT").
   - LƯU Ý ĐẶC BIỆT: TRƯỜNG HỌC KHÔNG ĐƯỢC ĐƯA VÀO ĐÂY! (Trường học phải đưa vào trường `school`).
   - Nếu trong mô tả chỉ có trường học mà không có công ty/nơi làm việc: để "Chưa cập nhật".

2. `Tên công ty chuẩn hóa (company_normalized)`:
   - Chuẩn hóa tên pháp nhân / loại hình doanh nghiệp đầy đủ theo chuẩn Việt Nam:
     * Viết tắt "JSC" -> "Công ty Cổ phần [Tên]" (Ví dụ: "GCW JSC" -> "Công ty Cổ phần GCW").
     * "Co., Ltd" / "Ltd" / "TNHH" -> "Công ty TNHH [Tên]" (Ví dụ: "iQIYI Vietnam" -> "Công ty TNHH iQIYI Vietnam", "Shopee Vietnam" -> "Công ty TNHH Shopee").
     * "Group" / "Tập đoàn" -> "Tập đoàn [Tên]" (Ví dụ: "Vingroup" -> "Tập đoàn Vingroup", "Sun Group" -> "Tập đoàn Sun Group").
     * "Corporation" / "Corp" -> "Tổng Công ty [Tên]" hoặc "Công ty Cổ phần [Tên]" (Ví dụ: "Viettel Post" -> "Tổng Công ty Cổ phần Bưu chính Viettel").
     * "Bank" / Ngân hàng -> "Ngân hàng TMCP [Tên]" (Ví dụ: "Vietcombank" -> "Ngân hàng TMCP Ngoại thương Việt Nam (Vietcombank)").
     * Cửa hàng / Spa / Thương hiệu cá nhân -> "Hộ kinh doanh / Thương hiệu [Tên]" (Ví dụ: "Món Ngon Bốn Mùa" -> "Thương hiệu Món Ngon Bốn Mùa").
     * Nếu không có công ty: để "Chưa cập nhật".

3. `Trường học / Học vấn (school)`:
   - Tên trường đại học, cao đẳng, học viện hoặc cơ sở giáo dục nếu có đề cập trong mô tả (Ví dụ: "Đại học Đà Lạt", "HUTECH - Đại học Công nghệ Tp.HCM", "Đại học Mở Hà Nội", "Trường Đại học Công nghệ Kỹ thuật TP. HCM", "Trường Đại học Văn hóa Hà Nội", "NEU").
   - Nếu không có thông tin trường học: để "Chưa cập nhật".

4. `Chức vụ (job_title)`:
   - Chức vụ cụ thể (VD: "Quản lý điều hành", "Giám đốc Marketing", "Chuyên viên", "Sinh viên").
   - Nếu đang học trường thì chức vụ ghi là "Sinh viên".
   - Nếu có nơi làm việc mà không ghi rõ chức vụ thì để "Nhân sự / Thành viên".
   - Nếu không có thông tin thì để "Chưa cập nhật".

5. `Cấp bậc (level)`:
   - "C-Level / Chủ DN", "Quản lý / Trưởng phòng", "Chuyên viên / Kỹ sư", "Freelance / Tự do", "Sinh viên / Học sinh", "Rác / Ảo", "Không xác định".

6. `Đánh giá (status)`:
   - "Hợp lệ", "Rác / Đùa cợt", "Không có thông tin việc làm".

7. `lead_tier`:
   - "Tiềm năng cao", "Tiềm năng trung bình", "Tiềm năng thấp", "Bỏ qua".

DANH SÁCH CẦN PHÂN TÍCH ({len(batch_items)} người):
"""
    for item in batch_items:
        prompt += f"- ID {item['id']}: Họ tên: \"{item['name']}\" | Mô tả Facebook: \"{item['text']}\"\n"

    prompt += """
YÊU CẦU ĐẦU RA:
Trả về DUY NHẤT một mảng JSON thuần túy (không kèm giải thích bên ngoài), theo cấu trúc:
[
  {
    "id": 1,
    "job_title": "Chức vụ cụ thể (hoặc Chưa cập nhật / Ảo / Đùa cợt)",
    "company": "Tên công ty / Doanh nghiệp / Tổ chức (hoặc Chưa cập nhật)",
    "company_normalized": "Tên công ty chuẩn hóa đầy đủ (hoặc Chưa cập nhật)",
    "school": "Tên trường đại học / học viện (hoặc Chưa cập nhật)",
    "level": "C-Level / Chủ DN | Quản lý / Trưởng phòng | Chuyên viên / Kỹ sư | Freelance / Tự do | Sinh viên / Học sinh | Rác / Ảo | Không xác định",
    "industry": "Lĩnh vực / Ngành nghề (hoặc Không xác định)",
    "status": "Hợp lệ | Rác / Đùa cợt | Không có thông tin việc làm",
    "lead_tier": "Tiềm năng cao | Tiềm năng trung bình | Tiềm năng thấp | Bỏ qua",
    "reason": "Giải thích ngắn gọn 1 câu lý do đánh giá"
  }
]
"""
    return prompt

def call_gemini_api(prompt: str, api_key: str, model: str = "gemini-2.5-flash-lite") -> List[Dict[str, Any]]:
    """
    Calls Google Gemini REST API directly with structured output format.
    Automatically handles rate limits (429) by switching to high-quota models like gemini-2.5-flash-lite.
    """
    headers = {"Content-Type": "application/json"}
    
    # Models to attempt in order if quota limit 429 is encountered
    candidate_models = [model]
    for fallback in ["gemini-2.5-flash-lite", "gemini-flash-latest"]:
        if fallback not in candidate_models:
            candidate_models.append(fallback)

    last_error = None

    for curr_model in candidate_models:
        url = f"https://generativelanguage.googleapis.com/v1beta/models/{curr_model}:generateContent?key={api_key}"
        payload = {
            "contents": [{
                "parts": [{"text": prompt}]
            }],
            "generationConfig": {
                "temperature": 0.1,
                "responseMimeType": "application/json"
            }
        }

        # Up to 2 retries per candidate model
        for attempt in range(2):
            try:
                response = requests.post(url, json=payload, headers=headers, timeout=60)
                if response.status_code == 200:
                    result = response.json()
                    text_content = result["candidates"][0]["content"]["parts"][0]["text"]
                    text_clean = text_content.strip()
                    if text_clean.startswith("```json"):
                        text_clean = text_clean[7:]
                    elif text_clean.startswith("```"):
                        text_clean = text_clean[3:]
                    if text_clean.endswith("```"):
                        text_clean = text_clean[:-3]
                    text_clean = text_clean.strip()
                    
                    try:
                        return json.loads(text_clean)
                    except json.JSONDecodeError:
                        m = re.search(r"\[\s*\{.*\}\s*\]", text_clean, re.DOTALL)
                        if m:
                            return json.loads(m.group(0))
                        raise
                elif response.status_code == 429:
                    # Rate limit exceeded: wait briefly and try fallback model if available
                    time.sleep(2 * (attempt + 1))
                    last_error = f"Gemini 429 Quota Exceeded on {curr_model}"
                    continue
                else:
                    last_error = f"Gemini API Error ({response.status_code}): {response.text}"
                    time.sleep(1.5)
            except Exception as e:
                last_error = str(e)
                time.sleep(1.5)

    if last_error:
        raise Exception(f"AI Batch Failed after fallback: {last_error}")
    return []

def call_openai_api(prompt: str, api_key: str, model: str = "gpt-4o-mini") -> List[Dict[str, Any]]:
    """
    Calls OpenAI / Compatible API for batch analysis.
    """
    url = "https://api.openai.com/v1/chat/completions"
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json"
    }
    payload = {
        "model": model,
        "messages": [
            {"role": "system", "content": "You are a professional B2B lead enrichment AI. Always output valid JSON array only."},
            {"role": "user", "content": prompt}
        ],
        "temperature": 0.1,
        "response_format": {"type": "json_object"}
    }
    response = requests.post(url, json=payload, headers=headers, timeout=60)
    if response.status_code == 200:
        content = response.json()["choices"][0]["message"]["content"]
        data = json.loads(content)
        if isinstance(data, list):
            return data
        elif isinstance(data, dict):
            for v in data.values():
                if isinstance(v, list):
                    return v
        return []
    else:
        raise Exception(f"OpenAI API Error ({response.status_code}): {response.text}")

def offline_heuristic_classify(name: str, text: str) -> Dict[str, Any]:
    """
    Offline fallback parser when no API key is provided or AI is unavailable.
    Extracts patterns using regex and dictionaries cleanly without raw metadata labels.
    """
    text_str = str(text) if pd.notna(text) else ""
    rule_res = rule_based_pre_classify(text_str)
    if rule_res:
        rule_res["company_normalized"] = rule_res.get("company", "Chưa cập nhật")
        rule_res["school"] = rule_res.get("school", "Chưa cập nhật")
        return rule_res

    # Clean text: remove deep crawl prefixes and metadata labels
    cleaned = text_str
    cleaned = re.sub(r"(?i)\b(Tiểu sử/Mô tả|Hạng mục/Lĩnh vực|Công việc|Học vấn|Sống tại|Liên hệ):\s*", "", cleaned)
    cleaned = re.sub(r"(?i)\bno workplaces to show\b", "", cleaned)
    cleaned = re.sub(r"(?i)\bno schools/universities to show\b", "", cleaned)
    cleaned = re.sub(r"(?i)\b\d+\s+bạn chung\b", "", cleaned)
    cleaned = re.sub(r"(?i)\b\d+\s+mutual friends\b", "", cleaned)
    cleaned = cleaned.strip(" ·|\n\r\t,")

    if not cleaned or len(cleaned) < 3:
        return {
            "job_title": "Chưa cập nhật",
            "company": "Chưa cập nhật",
            "company_normalized": "Chưa cập nhật",
            "school": "Chưa cập nhật",
            "level": "Không xác định",
            "industry": "Không rõ",
            "status": "Không có thông tin việc làm",
            "lead_tier": "Bỏ qua",
            "reason": "Chưa cài đặt thông tin việc làm công khai trên Facebook"
        }

    job_title = "Chưa cập nhật"
    company = "Chưa cập nhật"
    company_normalized = "Chưa cập nhật"
    school = "Chưa cập nhật"
    level = "Không xác định"
    status = "Hợp lệ"
    lead_tier = "Tiềm năng trung bình"
    industry = "Khác"

    cleaned_lower = cleaned.lower()

    # Detect School first
    school_match = re.search(r"((?:trường\s+)?đại\s+học[^\·\n,|]+|hutech[^\·\n,|]*|học\s+viện[^\·\n,|]+|cao\s+đẳng[^\·\n,|]+)", cleaned, re.IGNORECASE)
    if school_match:
        school = school_match.group(1).strip()

    # Split into segments by | or ·
    segments = [s.strip() for s in re.split(r"[·|]", cleaned) if s.strip()]

    found_job = None
    found_company = None

    for seg in segments:
        seg_clean = seg.strip()
        seg_lower = seg_clean.lower()
        if "no workplaces" in seg_lower or "no schools" in seg_lower or "bạn chung" in seg_lower:
            continue

        # Check for "Works at <Company>" or "Làm việc tại <Company>"
        m_works = re.match(r"^(?:Works at|Làm việc tại)\s+(.+)$", seg_clean, re.IGNORECASE)
        if m_works:
            c = m_works.group(1).strip()
            if not any(sk in c.lower() for sk in ["đại học", "cao đẳng", "học viện", "hutech", "neu"]):
                found_company = c
                if not found_job:
                    found_job = "Nhân sự / Thành viên"
            elif school == "Chưa cập nhật":
                school = c
            continue

        # Check for "<Job> at/tại <Company>"
        m_job_comp = re.search(r"^(.+?)\s+(?:tại|at)\s+(.+)$", seg_clean, re.IGNORECASE)
        if m_job_comp:
            cand_job = m_job_comp.group(1).strip()
            cand_comp = m_job_comp.group(2).strip()

            if re.match(r"^(?:Studied|Học|Đã học)\b", cand_job, re.IGNORECASE):
                if school == "Chưa cập nhật":
                    school = cand_comp
                continue

            if any(sk in cand_comp.lower() for sk in ["đại học", "cao đẳng", "học viện", "hutech", "neu"]):
                if school == "Chưa cập nhật":
                    school = cand_comp
            else:
                if not found_job or any(ck in cand_job.lower() for ck in ["ceo", "founder", "giám đốc", "director", "owner", "chủ", "co-founder"]):
                    found_job = cand_job
                    found_company = cand_comp
            continue

        if not found_job and not any(sk in seg_lower for sk in ["đại học", "cao đẳng", "học viện", "hutech", "neu", "studied", "học"]):
            if len(seg_clean) < 60:
                found_job = seg_clean

    if found_job:
        job_title = found_job
    if found_company:
        company = found_company

    # Normalize company name
    if company != "Chưa cập nhật":
        comp_clean = company
        if "JSC" in comp_clean:
            company_normalized = "Công ty Cổ phần " + comp_clean.replace("JSC", "").strip()
        elif any(k in comp_clean.upper() for k in ["LTD", "CO., LTD", "TNHH"]):
            company_normalized = "Công ty TNHH " + re.sub(r"(?i)\b(ltd|co\.,\s*ltd|tnhh)\b", "", comp_clean).strip()
        elif any(k in comp_clean.upper() for k in ["GROUP", "TẬP ĐOÀN"]):
            company_normalized = "Tập đoàn " + re.sub(r"(?i)\b(group|tập đoàn)\b", "", comp_clean).strip()
        elif "CORP" in comp_clean.upper():
            company_normalized = "Tổng Công ty " + re.sub(r"(?i)\bcorp(oration)?\b", "", comp_clean).strip()
        else:
            company_normalized = f"Doanh nghiệp / Đơn vị {comp_clean}"
    else:
        company_normalized = "Chưa cập nhật"

    # Level classification
    comb_lower = (job_title + " " + cleaned_lower).lower()
    if any(k in comb_lower for k in ["ceo", "founder", "sáng lập", "tổng giám đốc", "chủ tịch", "chủ chuỗi", "chủ shop", "chủ", "co-founder"]):
        level = "C-Level / Chủ DN"
        lead_tier = "Tiềm năng cao"
    elif any(k in comb_lower for k in ["giám đốc", "trưởng phòng", "head of", "quản lý", "trưởng nhóm", "team lead", "director", "manager"]):
        level = "Quản lý / Trưởng phòng"
        lead_tier = "Tiềm năng cao"
    elif any(k in comb_lower for k in ["kỹ sư", "chuyên viên", "bác sĩ", "luật sư", "senior", "developer", "architect", "expert", "chuyên gia"]):
        level = "Chuyên viên / Kỹ sư"
        lead_tier = "Tiềm năng trung bình"
    elif school != "Chưa cập nhật" and company == "Chưa cập nhật" and job_title in ["Chưa cập nhật", "Sinh viên"]:
        level = "Sinh viên / Học sinh"
        lead_tier = "Tiềm năng thấp"
        job_title = "Sinh viên"
        status = "Hợp lệ"
    elif any(k in comb_lower for k in ["freelance", "tự do"]):
        level = "Freelance / Tự do"
        lead_tier = "Tiềm năng thấp"
    else:
        if job_title == "Chưa cập nhật" and company == "Chưa cập nhật" and school == "Chưa cập nhật":
            status = "Không có thông tin việc làm"
            lead_tier = "Bỏ qua"

    # Industry detection
    if any(k in comb_lower for k in ["bất động sản", "real estate", "đất xanh", "sun group", "vingroup", "nhà đất"]):
        industry = "Bất động sản"
    elif any(k in comb_lower for k in ["công nghệ", "tech", "fpt", "software", "solution", "it", "gcw", "ai"]):
        industry = "Công nghệ thông tin"
    elif any(k in comb_lower for k in ["cà phê", "coffee", "f&b", "ẩm thực", "nhà hàng", "món ngon", "wines"]):
        industry = "F&B / Ẩm thực"
    elif any(k in comb_lower for k in ["marketing", "media", "agency", "quảng cáo", "communications", "training"]):
        industry = "Marketing & Truyền thông"
    elif any(k in comb_lower for k in ["chứng khoán", "ngân hàng", "bank", "tài chính", "đầu tư", "tax", "thuế"]):
        industry = "Tài chính / Ngân hàng"
    elif any(k in comb_lower for k in ["thời trang", "boutique", "spa", "thẩm mỹ", "viona", "luxury"]):
        industry = "Dịch vụ & Làm đẹp"
    elif any(k in comb_lower for k in ["logistics", "transport", "vận tải", "giao nhận"]):
        industry = "Vận tải / Logistics"
    elif any(k in comb_lower for k in ["build", "design", "thiết kế", "kiến trúc", "xây dựng", "thước tầm"]):
        industry = "Xây dựng / Kiến trúc"

    return {
        "job_title": job_title,
        "company": company,
        "company_normalized": company_normalized,
        "school": school,
        "level": level,
        "industry": industry,
        "status": status,
        "lead_tier": lead_tier,
        "reason": "Phân tích tự động bằng bộ quy tắc thông minh (Offline)"
    }

def process_friends_dataframe(
    df: pd.DataFrame,
    name_col: str,
    link_col: str,
    info_cols: Any,
    api_key: Optional[str] = None,
    api_provider: str = "gemini",
    model_name: str = "gemini-2.5-flash",
    target_criteria: str = "",
    batch_size: int = 25,
    progress_callback=None
) -> pd.DataFrame:
    """
    Main pipeline to process rows in batches, applying AI classification with fallback.
    """
    total_rows = len(df)
    results = []

    # Handle single or multiple info cols
    if isinstance(info_cols, str):
        info_col_list = [info_cols]
    elif isinstance(info_cols, (list, tuple)):
        info_col_list = list(info_cols)
    else:
        info_col_list = []

    # Prepare standard items
    prepared_items = []
    for idx, row in df.iterrows():
        name_val = str(row.get(name_col, "")).strip() if pd.notna(row.get(name_col)) else ""
        link_val = str(row.get(link_col, "")).strip() if pd.notna(row.get(link_col)) else ""
        
        # Combine all info columns together!
        info_parts = []
        for ic in info_col_list:
            if ic in df.columns and pd.notna(row.get(ic)):
                v = str(row.get(ic, "")).strip()
                if v and v.lower() != "nan" and not v.startswith("http"):
                    info_parts.append(v)
        full_info = " · ".join(info_parts)
        
        prepared_items.append({
            "id": idx + 1,
            "name": name_val,
            "link": link_val,
            "text": full_info,
            "orig_row": row.to_dict()
        })

    # Process batch by batch
    for start_idx in range(0, total_rows, batch_size):
        batch = prepared_items[start_idx : start_idx + batch_size]
        
        # If API key is not provided, use offline rule classifier
        if not api_key:
            for item in batch:
                cls_res = offline_heuristic_classify(item["name"], item["text"])
                results.append({
                    "Họ và tên": item["name"],
                    "Chức vụ": cls_res["job_title"],
                    "Tên công ty / Đơn vị": cls_res["company"],
                    "Tên công ty chuẩn hóa": cls_res.get("company_normalized", "Chưa cập nhật"),
                    "Trường học / Học vấn": cls_res.get("school", "Chưa cập nhật"),
                    "Cấp bậc": cls_res["level"],
                    "Lĩnh vực / Ngành nghề": cls_res["industry"],
                    "Đánh giá": cls_res["status"],
                    "Phân loại Lead": cls_res["lead_tier"],
                    "Link Facebook": item["link"],
                    "Mô tả gốc": item["text"],
                    "Ghi chú AI": cls_res["reason"]
                })
        else:
            # Need to call AI API
            # First check for quick pre-filters (jokes, mutual friends only, blank) to save tokens & time
            ai_batch = []
            id_to_prefilter = {}
            for item in batch:
                pre = rule_based_pre_classify(item["text"])
                if pre and pre["status"] in ["Rác / Đùa cợt", "Không có thông tin việc làm", "Không có thông tin"]:
                    id_to_prefilter[item["id"]] = pre
                else:
                    ai_batch.append(item)

            ai_results_map = {}
            if ai_batch:
                prompt = build_gemini_prompt(ai_batch, target_criteria)
                try:
                    if api_provider == "gemini":
                        ai_resp = call_gemini_api(prompt, api_key, model=model_name)
                    else:
                        ai_resp = call_openai_api(prompt, api_key, model=model_name)

                    for r in ai_resp:
                        if isinstance(r, dict) and "id" in r:
                            ai_results_map[int(r["id"])] = r
                except Exception as e:
                    # If AI fails, fallback to offline rules for this batch so processing continues
                    print(f"Warning: AI batch error ({e}), falling back to offline rules for batch {start_idx}...")
                    for item in ai_batch:
                        ai_results_map[item["id"]] = offline_heuristic_classify(item["name"], item["text"])

            # Merge results for this batch
            for item in batch:
                item_id = item["id"]
                if item_id in id_to_prefilter:
                    res_data = id_to_prefilter[item_id]
                elif item_id in ai_results_map:
                    res_data = ai_results_map[item_id]
                else:
                    res_data = offline_heuristic_classify(item["name"], item["text"])

                results.append({
                    "Họ và tên": item["name"],
                    "Chức vụ": res_data.get("job_title", "Chưa cập nhật"),
                    "Tên công ty / Đơn vị": res_data.get("company", "Chưa cập nhật"),
                    "Tên công ty chuẩn hóa": res_data.get("company_normalized", "Chưa cập nhật"),
                    "Trường học / Học vấn": res_data.get("school", "Chưa cập nhật"),
                    "Cấp bậc": res_data.get("level", "Không xác định"),
                    "Lĩnh vực / Ngành nghề": res_data.get("industry", "Không rõ"),
                    "Đánh giá": res_data.get("status", "Hợp lệ"),
                    "Phân loại Lead": res_data.get("lead_tier", "Tiềm năng trung bình"),
                    "Link Facebook": item["link"],
                    "Mô tả gốc": item["text"],
                    "Ghi chú AI": res_data.get("reason", "")
                })

            # Small polite pause to respect rate limits only when AI was actually called
            if ai_batch:
                time.sleep(1.0)

        # Notify progress callback
        if progress_callback:
            current_done = min(start_idx + batch_size, total_rows)
            progress_callback(current_done, total_rows)

    return pd.DataFrame(results)

def export_styled_excel(df: pd.DataFrame, output_path: Any = None) -> Any:
    """
    Exports processed dataframe into a professionally styled Excel file.
    Can write to a file path, a file-like buffer (io.BytesIO), or return raw bytes if output_path is None.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Khách Hàng Tiềm Năng"
    ws.views.sheetView[0].showGridLines = True

    # Color definitions
    header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid") # Dark Navy
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    
    tier_high_fill = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid") # Soft green
    tier_mid_fill = PatternFill(start_color="CFF4FC", end_color="CFF4FC", fill_type="solid")  # Soft blue
    tier_trash_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")# Soft red
    tier_none_fill = PatternFill(start_color="E2E3E5", end_color="E2E3E5", fill_type="solid") # Soft gray

    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    # Headers
    headers = list(df.columns)
    ws.append(headers)

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    # Rows
    for row_idx, row_data in enumerate(df.values, 2):
        ws.append(list(row_data))
        ws.row_dimensions[row_idx].height = 22
        
        tier_val = str(row_data[headers.index("Phân loại Lead")]) if "Phân loại Lead" in headers else ""
        status_val = str(row_data[headers.index("Đánh giá")]) if "Đánh giá" in headers else ""

        # Row color highlight based on lead tier
        row_fill = None
        if "Tiềm năng cao" in tier_val or "C-Level" in str(row_data):
            row_fill = tier_high_fill
        elif "Tiềm năng trung bình" in tier_val:
            row_fill = tier_mid_fill
        elif "Rác" in status_val or "Bỏ qua" in tier_val:
            row_fill = tier_trash_fill
        elif "Không có thông tin" in status_val:
            row_fill = tier_none_fill

        has_li_link = False
        if "Link LinkedIn" in headers:
            li_idx = headers.index("Link LinkedIn")
            if "linkedin.com/in" in str(row_data[li_idx]):
                has_li_link = True

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            header_name = headers[col_idx - 1]

            if has_li_link:
                cell.font = Font(name="Arial", size=10, bold=True)
            else:
                cell.font = Font(name="Arial", size=10)

            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

            # Make Facebook link clickable
            if header_name == "Link Facebook" and str(cell.value).startswith("http"):
                cell.hyperlink = str(cell.value)
                cell.font = Font(name="Arial", size=10, color="0000FF", underline="single", bold=has_li_link)

            # Make LinkedIn link clickable & highlighted
            if header_name == "Link LinkedIn" and "linkedin.com/in" in str(cell.value):
                cell.hyperlink = str(cell.value)
                cell.font = Font(name="Arial", size=10, color="0A66C2", underline="single", bold=True)
                cell.fill = PatternFill(start_color="D1E7DD", end_color="D1E7DD", fill_type="solid")

            # Apply background tint to key columns (Phân loại Lead, Đánh giá)
            if header_name in ["Phân loại Lead", "Cấp bậc", "Đánh giá"] and row_fill:
                cell.fill = row_fill

    # Auto fit column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    if output_path is None:
        import io
        bio = io.BytesIO()
        wb.save(bio)
        bio.seek(0)
        return bio.getvalue()
    elif hasattr(output_path, "write"):
        wb.save(output_path)
        return output_path
    else:
        wb.save(output_path)
        return output_path
